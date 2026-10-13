"""Extração auditada e reproduzível das bases MES e SAP a partir da planilha original."""
from pathlib import Path
from collections import Counter, defaultdict
from datetime import datetime, time, timedelta
import csv
import hashlib
import json
import re
from statistics import median
import unicodedata

from openpyxl import load_workbook
from openpyxl.utils.datetime import from_excel

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "data" / "extracoes"
FILE = ROOT / "fontes" / "Base de dados Moura - COMPET - Manutencao Inteligente.xlsx"


def txt(v):
    return re.sub(r"\s+", " ", str(v or "")).strip()


def serial(v):
    return v.isoformat() if isinstance(v, (datetime, time)) else v


def date(v, epoch):
    if v is None: return None
    if isinstance(v, datetime): return v
    if isinstance(v, (int, float)): return from_excel(v, epoch)
    try: return datetime.fromisoformat(str(v))
    except ValueError: return None


def dt(day, clock, epoch):
    day = date(day, epoch)
    if not isinstance(day, datetime): return None
    if isinstance(clock, datetime): clock = clock.time()
    if isinstance(clock, (int, float)):
        clock = (datetime.min + timedelta(days=clock)).time()
    if isinstance(clock, str):
        try: clock = time.fromisoformat(clock)
        except ValueError: return None
    if not isinstance(clock, time): return None
    return datetime.combine(day.date(), clock)


def dump(name, value):
    (OUT / name).write_text(json.dumps(value, ensure_ascii=False, indent=2, default=serial), encoding="utf-8")


def write_csv(name, rows):
    keys = list(dict.fromkeys(k for row in rows for k in row))
    with (OUT / name).open("w", encoding="utf-8-sig", newline="") as f:
        writer = csv.DictWriter(f, fieldnames=keys)
        writer.writeheader(); writer.writerows(rows)


def main():
    OUT.mkdir(parents=True, exist_ok=True)
    digest = hashlib.sha256(FILE.read_bytes()).hexdigest()
    w = load_workbook(FILE, read_only=True, data_only=True)
    formula_w = load_workbook(FILE, read_only=True, data_only=False)
    inventory = {"arquivo": FILE.name, "sha256": digest, "excel_epoch": w.epoch.isoformat(),
        "abas": []}
    tables = {}
    quality = {}
    dictionary = []
    for sheet in w:
        headers = [str(v) if v is not None else f"coluna_{i+1}" for i,v in enumerate(next(sheet.iter_rows(values_only=True)))]
        rows = []
        for line, values in enumerate(sheet.iter_rows(min_row=2, values_only=True), 2):
            if not any(v is not None for v in values): continue
            rows.append({"_arquivo": FILE.name, "_aba": sheet.title, "_linha_excel": line,
                "_id_auditoria": ("MES" if sheet.title.startswith("MES") else "SAP") + f"-R{line:04d}",
                **{key: serial(value) for key,value in zip(headers, values)}})
        formulas = [{"celula": c.coordinate, "formula": c.value} for r in formula_w[sheet.title] for c in r if c.data_type == "f"]
        errors = [{"celula": c.coordinate, "valor": c.value} for r in sheet for c in r if c.data_type == "e"]
        columns = {}
        for h in headers:
            vals = [r[h] for r in rows]
            missing = sum(v is None or (isinstance(v,str) and not v.strip()) for v in vals)
            counter = Counter(str(v) for v in vals if v is not None)
            columns[h] = {"ausentes": missing, "denominador": len(rows), "distintos": len(counter),
                "tipos_exportados": dict(Counter(type(v).__name__ for v in vals)), "exemplo": next((v for v in vals if v is not None),None)}
            dictionary.append(f"| {sheet.title} | `{h}` | {','.join(columns[h]['tipos_exportados'])} | {missing}/{len(rows)} | `{str(columns[h]['exemplo'])[:95]}` |")
        full = Counter(tuple(str(r[h]) for h in headers) for r in rows)
        duplicates = [[r['_id_auditoria'] for r in rows if tuple(str(r[h]) for h in headers)==k]
                      for k,n in full.items() if n>1]
        quality[sheet.title] = {"registros": len(rows), "colunas": columns, "grupos_duplicatas_exatas": duplicates,
            "formulas": len(formulas), "erros_excel": errors}
        tables[sheet.title] = rows
        inventory['abas'].append({"aba": sheet.title, "cabecalho_linha":1, "registros": len(rows), "cabecalhos_originais":headers,
            "csv": "MES.csv" if sheet.title.startswith("MES") else "SAP.csv", "formulas":formulas})
    mes = tables['MES - PRODUÇÃO']; sap = tables['SAP - MANUTENÇÃO']
    mes_issues = defaultdict(list); mes_groups=Counter()
    for r in mes:
        r['maquina_norm'] = txt(r['MÁQUINA ']); r['equipamento_norm'] = txt(r['EQUIPAMENTO'])
        r['componente_norm'] = txt(r['Componente']); r['familia_norm'] = txt(r['Famila'])
        start=date(r['Inicio'], w.epoch); end=date(r['Término'], w.epoch); day=date(r['Data'],w.epoch)
        r['inicio_iso']=serial(start);r['fim_iso']=serial(end)
        r['duracao_derivada_h']=(end-start).total_seconds()/3600 if start and end else None
        if start and day and start.date()!=day.date():mes_issues['data_diferente_inicio'].append(r['_id_auditoria'])
        if start and end and start.date()!=end.date():mes_issues['vira_dia'].append(r['_id_auditoria'])
        if r['duracao_derivada_h'] is not None and r['duracao_derivada_h']<0:mes_issues['duracao_negativa'].append(r['_id_auditoria'])
        if r['Horas'] is None:mes_issues['horas_ausentes'].append(r['_id_auditoria'])
        elif r['duracao_derivada_h'] is not None and abs(float(r['Horas'])-r['duracao_derivada_h'])*3600>1:mes_issues['diferenca_duracao_mais_1s'].append(r['_id_auditoria'])
        if r['Componente'] != r['componente_norm']:mes_issues['espacos_componente'].append(r['_id_auditoria'])
        mes_groups[(r['maquina_norm'],r['familia_norm'],r['equipamento_norm'],r['componente_norm'])]+=1
    sap_issues=defaultdict(list)
    for r in sap:
        location=txt(r['Loc.instalação']);match=re.search(r'\.MAQ([A-D])(?:\.|$)',location)
        r['maquina_inferida']=match.group(1) if match else None
        r['mapeamento_maquina_status']='inferido de segmento MAQ; validar tecnicamente'
        start=dt(r['Dta.início real'],r['InícReal hr.'],w.epoch); end=dt(r['Dt.real fim'],r['Fim real'],w.epoch)
        base=date(r['Data-base iníc.'],w.epoch)
        r['inicio_iso']=serial(start);r['fim_iso']=serial(end);r['data_base_iso']=serial(base)
        r['duracao_intervalo_registrado_h']=(end-start).total_seconds()/3600 if start and end else None
        if not start or not end:sap_issues['inicio_ou_fim_invalido'].append(r['_id_auditoria'])
        if r['duracao_intervalo_registrado_h'] is not None and r['duracao_intervalo_registrado_h']<0:sap_issues['intervalo_negativo'].append(r['_id_auditoria'])
        if start and end and start.date()!=end.date():sap_issues['vira_dia'].append(r['_id_auditoria'])
        if start and base and start.date()!=base.date():sap_issues['data_base_diferente_inicio'].append(r['_id_auditoria'])
        if not match:sap_issues['maquina_nao_mapeada'].append(r['_id_auditoria'])
    overlap=[]
    for machine in sorted({r['maquina_norm'] for r in mes}):
        group=sorted([r for r in mes if r['maquina_norm']==machine],key=lambda r:r['inicio_iso'] or '')
        active=[]
        for row in group:
            start=row['inicio_iso'];active=[a for a in active if a['fim_iso'] and start and a['fim_iso']>start]
            for a in active:overlap.append({'maquina':machine,'id_1':a['_id_auditoria'],'id_2':row['_id_auditoria']})
            active.append(row)
    quality['MES - PRODUÇÃO']['verificacoes_temporais_e_textuais']=dict(mes_issues)
    quality['MES - PRODUÇÃO']['pares_sobrepostos_mesma_maquina']=overlap
    quality['SAP - MANUTENÇÃO']['verificacoes_temporais']=dict(sap_issues)
    orders=Counter(str(r['N. Ordem']) for r in sap)
    quality['SAP - MANUTENÇÃO']['ordens_repetidas']={k:n for k,n in orders.items() if n>1}
    summary = {'fonte_sha256':digest,
        'registros':{'MES':len(mes),'SAP':len(sap)},
        'MES_maquinas':dict(Counter(r['maquina_norm'] for r in mes)),
        'SAP_maquinas_inferidas':dict(Counter(r['maquina_inferida'] for r in sap)),
        'MES_familias':dict(Counter(r['familia_norm'] for r in mes)),
        'SAP_tipos':dict(Counter(r['Tipo'] for r in sap)),
        'SAP_atividades':dict(Counter(str(r['Tip.ativ.manut.']) for r in sap)),
        'MES_problemas':{k:len(v) for k,v in mes_issues.items()},
        'SAP_problemas':{k:len(v) for k,v in sap_issues.items()},
        'MES_pares_sobrepostos':dict(Counter(r['maquina'] for r in overlap)),
        'cobertura':{'MES':[min(r['inicio_iso'] for r in mes),max(r['fim_iso'] for r in mes)],
            'SAP':[min(r['inicio_iso'] for r in sap if r['inicio_iso']),max(r['fim_iso'] for r in sap if r['fim_iso'])]},
        'recorrencias_A': [{'maquina':a,'familia':b,'equipamento':c,'componente':d,'contagem':n}
            for (a,b,c,d),n in mes_groups.most_common() if a=='A'],
        'limites':['maquina SAP inferida, nao confirmada tecnicamente','códigos de componente sem dicionario tecnico',
            'exposicao/horas operadas, marcos de chamados, custos e HH nao presentes nas colunas',
            'duracao SAP e intervalo entre marcos cadastrados, nao tempo de reparo validado']}
    sample = sorted([r for r in sap if r['maquina_inferida']=='A'],key=lambda r:r['_linha_excel'])
    # Amostra deterministica de inspeção, incluindo diversidade de tipos/atividades; nao e verdade rotulada.
    buckets=defaultdict(list)
    for r in sample:buckets[(r['Tipo'],r['Tip.ativ.manut.'])].append(r)
    selected=[]
    for rows in buckets.values():
        for index in sorted({0,len(rows)//2,len(rows)-1}):selected.append(rows[index])
    selected=sorted({r['_id_auditoria']:r for r in selected}.values(),key=lambda r:r['_linha_excel'])
    write_csv('MES.csv',mes);write_csv('SAP.csv',sap);write_csv('AMOSTRA_SAP_A.csv',selected)
    dump('inventario.json',inventory);dump('perfil_qualidade.json',quality);dump('resumo.json',summary)
    dictionary_text='# Dicionário observado dos dados\n\nTipos referem-se a exportação; datas SAP numéricas são seriais Excel. Linhas originais e cabeçalhos preservados.\n\n'
    dictionary_text+='| Aba | Campo original | Tipos exportados | Ausentes/total | Exemplo |\n|---|---|---|---|---|\n'+'\n'.join(dictionary)
    dictionary_text+='\n\nMES: Horas é duração informada da ocorrência, comparada com fim–início; Data pode representar dia operacional distinto da data do início. EQUIPAMENTO é a subdivisão documentada no enunciado; códigos de componente exigem dicionário técnico.\n\nSAP: datas numéricas convertidas com epoch do arquivo; intervalo início–fim não é HH ou tempo de reparo confirmado. N. Ordem identifica uma ordem, cujo número de linhas é auditado. MAQA etc. sugere máquina, com mapeamento inferido. Tipo Corretiva Planejada não é convertido em Preventiva. Não há marcos de solicitação/aceite, exposição operacional, custos ou conteúdo de plano preventivo nas colunas.\n'
    (OUT/'DICIONARIO_DADOS.md').write_text(dictionary_text,encoding='utf-8')
    (OUT/'DIAGNOSTICO_BASES.md').write_text('# Diagnóstico reproduzível\n\nEvidências numéricas em resumo.json e perfil_qualidade.json.\n\n```json\n'+json.dumps(summary,ensure_ascii=False,indent=2)+'\n```\n\nNão foram excluídas linhas nem corrigidos horários. Normalizações são derivadas e conservam as colunas originais. Detecção de sobreposição usa intervalos da mesma máquina, sem declarar erro ou somar durações como perda exclusiva.\n',encoding='utf-8')
    w.close();formula_w.close()
    assert hashlib.sha256(FILE.read_bytes()).hexdigest()==digest
    print(json.dumps(summary,ensure_ascii=False,indent=2))


if __name__=='__main__':main()
